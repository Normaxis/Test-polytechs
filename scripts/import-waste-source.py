import openpyxl, json, datetime, re, hashlib, sys, collections
input_path, output_path = sys.argv[1:]
w = openpyxl.load_workbook(input_path, data_only=True)
s = w['enlèvements']
def val(v):
    if isinstance(v, (datetime.datetime, datetime.date)): return v.isoformat()[:10]
    return '' if v is None else str(v).strip()
def num(v):
    return float(v) if isinstance(v,(int,float)) and not isinstance(v,bool) else None
headers=[val(c.value) for c in s[5]][:59]
flags={4:'DIB',5:'Déchets clients',6:'Purges',7:'Cartons',8:'Plastiques',9:'Terre et granulés',10:'4 flux',11:'Big-bags',12:'Plastiques',13:'Plastiques',14:'Plastiques',15:'Plastiques',16:'Déchets dangereux',17:'Autres',18:'Ferraille',19:'Palettes',20:'Bois',21:'Déchets ménagers',22:'Papier',23:'Retraits salariés',24:'Intervenants extérieurs',25:'Déchetterie'}
records=[]
for row in s.iter_rows(min_row=6):
    v=[c.value for c in row]; n=row[0].row
    if not v[29]:continue
    raw={openpyxl.utils.get_column_letter(i+1):val(x) for i,x in enumerate(v[:59]) if x is not None}
    original=val(v[0]); planned=bool(re.search('prév|demande|planification',original,re.I))
    date='';issues=[]
    if isinstance(v[0],(datetime.datetime,datetime.date)):date=val(v[0])
    else:
        match=re.search(r'(?<!\d)(\d{1,2})/(\d{1,2})/(\d{4}|\d{2})(?!\d)',original)
        if match:
            day,month,year=map(int,match.groups());year=year+2000 if year<100 else year
            try:date=datetime.date(year,month,day).isoformat()
            except ValueError:pass
    if not date:issues.append('Date absente ou à corriger')
    q=num(v[35]); estimated=num(v[34]);unit='t'
    if v[19] or val(v[29]).lower()=='palettes':unit='unité'
    if q is not None and unit=='t' and q>100:
        unit='à vérifier';issues.append('Unité de la quantité à confirmer')
    if q is None and not planned:issues.append('Quantité réelle non renseignée')
    danger={'OUI':'Oui','NON':'Non'}.get(val(v[33]).upper(),'À vérifier')
    if danger=='À vérifier':issues.append('Dangerosité non renseignée')
    if '*' in val(v[30]) and danger=='Non':issues.append('Code étoilé et dangerosité NON dans la source')
    chosen=[k for k in flags if val(v[k]).upper().startswith('X')]
    family=flags.get(next((k for k in chosen if k not in [4,5,16,17,23,24,25]),chosen[0] if chosen else 17),'Autres')
    excluded=bool(v[26])
    if danger=='Oui' and not v[28]:issues.append('Référence BSD à renseigner')
    records.append(dict(id='excel-dechets-'+str(n),revision=0,date=date,dateSource=original,status='Prévu' if planned else 'Enlevé',type=val(v[29]),family=family,code=val(v[30]),danger=danger,quantity=q,estimate=estimated,unit=unit,excluded=excluded,client=bool(v[5]),collector=val(v[43]).split('\n')[0],destination=val(v[49]),treatment=val(v[42]),bsd=val(v[28]),owner='',notes='',documentUrl='',sourceRow=n,issues=issues,raw=raw))
reference=[]
for row in w['désignation ADR'].iter_rows(min_row=5):
    if any(c.value is not None for c in row): reference.append({'row':row[0].row,'raw':{openpyxl.utils.get_column_letter(i+1):val(c.value) for i,c in enumerate(row) if c.value is not None}})
out={'source':'Imp2.1305b_Suivi enlèvements déchets.xlsx','sha256':hashlib.sha256(open(input_path,'rb').read()).hexdigest(),'headers':headers,'records':records,'adrReference':reference,'treatments':[str(r[0]) for r in list(w['Données'].values)[1:] if r[0]]}
open(output_path,'w').write(json.dumps(out,ensure_ascii=False,separators=(',',':')))
print(json.dumps({'records':len(records),'years':dict(collections.Counter(r['date'][:4] or 'sans date' for r in records)),'planned':sum(r['status']=='Prévu' for r in records),'issues':sum(bool(r['issues']) for r in records),'units':dict(collections.Counter(r['unit'] for r in records)),'reference':len(reference)},ensure_ascii=False))
